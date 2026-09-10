import json
import subprocess
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

output_dir = Path(__file__).resolve().parent
remote_script = r'''
require("dotenv").config({ quiet: true });
const mongoose = require("mongoose");
const LicenseRecord = require("./src/models/licenseRecord");
(async () => {
  await mongoose.connect(process.env.MONGO_URI || process.env.MONGODB_URI);
  const rows = await LicenseRecord.aggregate([
    { $set: { normalizedLicense: { $toUpper: { $trim: { input: { $ifNull: ["$license_number", ""] } } } } } },
    { $match: { normalizedLicense: { $ne: "" } } },
    { $setWindowFields: { partitionBy: "$normalizedLicense", output: { duplicateCount: { $count: {} } } } },
    { $match: { duplicateCount: { $gt: 1 } } },
    { $sort: { normalizedLicense: 1, business_name: 1 } },
    { $project: { normalizedLicense: 1, duplicateCount: 1, business_name: 1, dba: 1, business_address: 1, city: 1, stateName: 1, postal_code: 1, license_status: 1, license_type: 1, expiration_date: 1, claimed: 1, canojaVerified: 1, contact_information: 1, government_source: 1, sourceType: 1, createdAt: 1 } }
  ]);
  process.stdout.write(JSON.stringify(rows));
  await mongoose.disconnect();
})().catch(async (error) => { console.error(error.message); await mongoose.disconnect().catch(() => {}); process.exit(1); });
'''
result = subprocess.run([
    "ssh", "-i", "/Users/test/Downloads/canoja-new.pem", "-o", "BatchMode=yes",
    "-o", "ConnectTimeout=10", "ubuntu@54.227.140.191",
    "cd /home/ubuntu/workspace/server && source /home/ubuntu/.nvm/nvm.sh && node -",
], input=remote_script, text=True, capture_output=True, check=True)
rows = json.loads(result.stdout)
group_count = len({row["normalizedLicense"] for row in rows})

workbook = Workbook()
sheet = workbook.active
sheet.title = "Duplicate licenses"
sheet.sheet_view.showGridLines = False
sheet["A2"] = "Production duplicate license records"
sheet["A2"].font = Font(name="Arial", size=16, bold=True, color="18352A")
sheet["A4"], sheet["B4"] = "Production records", 28969
sheet["D4"], sheet["E4"] = "Duplicate licenses", group_count
sheet["G4"], sheet["H4"] = "Affected records", len(rows)
for cell in sheet[4]:
    cell.font = Font(name="Arial", size=10, bold=True, color="18352A")
sheet["A5"] = "License numbers are normalized using trim and uppercase. Blank license numbers are excluded."
sheet["A5"].font = Font(name="Arial", size=10, italic=True, color="617182")

headers = ["License number", "Duplicate count", "Record ID", "Business name", "DBA", "Address", "City", "State", "Postal code", "License status", "License type", "Expiration date", "Claimed", "Canoja verified", "Phone", "Email", "Website", "Source provider", "Source URL", "Created at"]
for column, header in enumerate(headers, 1):
    cell = sheet.cell(7, column, header)
    cell.fill = PatternFill("solid", fgColor="1B6B46")
    cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def parsed_date(value):
    if not value:
        return None
    return datetime.fromisoformat(value.replace("Z", "+00:00")).replace(tzinfo=None)

for index, row in enumerate(rows, 8):
    contact = row.get("contact_information") or {}
    source = row.get("government_source") or {}
    values = [
        row["normalizedLicense"], row["duplicateCount"], str(row["_id"]), row.get("business_name", ""),
        row.get("dba", ""), row.get("business_address", ""), row.get("city", ""), row.get("stateName", ""),
        str(row.get("postal_code", "") or ""), row.get("license_status", ""), row.get("license_type", ""),
        parsed_date(row.get("expiration_date")), bool(row.get("claimed")), bool(row.get("canojaVerified")),
        contact.get("phone", ""), contact.get("email", ""), contact.get("website", ""),
        source.get("provider") or row.get("sourceType", ""), source.get("url", ""), parsed_date(row.get("createdAt")),
    ]
    for column, value in enumerate(values, 1):
        cell = sheet.cell(index, column, value)
        cell.font = Font(name="Arial", size=10, color="18212B")
        cell.alignment = Alignment(vertical="center")
        cell.border = Border(bottom=Side(style="thin", color="DCE7E1"))

last_row = len(rows) + 7
sheet.freeze_panes = "D8"
sheet.auto_filter.ref = f"A7:T{last_row}"
for row in range(8, last_row + 1):
    sheet.cell(row, 12).number_format = "mm/dd/yyyy"
    sheet.cell(row, 20).number_format = "mm/dd/yyyy hh:mm"
sheet.conditional_formatting.add(f"N8:N{last_row}", CellIsRule(operator="equal", formula=["TRUE"], fill=PatternFill("solid", fgColor="E7F7EE"), font=Font(bold=True, color="1B6B46")))
widths = [18, 13, 25, 34, 28, 42, 20, 18, 13, 16, 28, 15, 11, 16, 18, 28, 34, 20, 42, 20]
for column, width in enumerate(widths, 1):
    sheet.column_dimensions[sheet.cell(7, column).column_letter].width = width
sheet.row_dimensions[7].height = 30

output_path = output_dir / "production_duplicate_license_records.xlsx"
workbook.save(output_path)
check = load_workbook(output_path, read_only=True, data_only=False)
check_sheet = check["Duplicate licenses"]
assert check_sheet.max_row == last_row
assert check_sheet.max_column == 20
assert check_sheet["E4"].value == group_count
assert check_sheet["H4"].value == len(rows)
assert all(check_sheet.cell(row, 1).value for row in range(8, last_row + 1))
print(json.dumps({"path": str(output_path), "duplicateLicenses": group_count, "affectedRecords": len(rows), "rows": last_row}))
