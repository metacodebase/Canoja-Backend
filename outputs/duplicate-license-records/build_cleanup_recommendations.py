import json
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

base_dir = Path(__file__).resolve().parent
rows = json.loads((base_dir / "duplicate_license_research.json").read_text())
groups = {}
for row in rows:
    groups.setdefault(row["license"], []).append(row)

decision_counts = Counter(group[0]["decision"] for group in groups.values())
keep_groups = sum(decision == "KEEP AND MERGE DUPLICATES" for decision in (group[0]["decision"] for group in groups.values()))
manual_groups = len(groups) - keep_groups
remove_candidates = sum(row["recommendation"] == "MERGE INTO CANONICAL, THEN REMOVE" for row in rows)

wb = Workbook()
summary = wb.active
summary.title = "Summary"
groups_sheet = wb.create_sheet("Group decisions")
records_sheet = wb.create_sheet("Record recommendations")
sources_sheet = wb.create_sheet("Sources")

dark = "18352A"
green = "1B6B46"
light_green = "E7F7EE"
amber = "FFF4CC"
red = "FDE8E7"
gray = "617182"
line = Side(style="thin", color="DCE7E1")

def title(sheet, text):
    sheet.sheet_view.showGridLines = False
    sheet["A2"] = text
    sheet["A2"].font = Font(name="Arial", size=16, bold=True, color=dark)
    sheet["A3"].border = Border(bottom=Side(style="thin", color="2DA96D"))

def write_header(sheet, row_number, headers):
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(row_number, column, header)
        cell.fill = PatternFill("solid", fgColor=green)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[row_number].height = 30

def style_body(sheet, start_row, end_row, end_column):
    for row_number in range(start_row, end_row + 1):
        for column in range(1, end_column + 1):
            cell = sheet.cell(row_number, column)
            cell.font = Font(name="Arial", size=10, color="18212B")
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(bottom=line)

title(summary, "Duplicate license cleanup recommendations")
summary["A4"], summary["B4"] = "Duplicate license groups", len(groups)
summary["A5"], summary["B5"] = "Affected records", len(rows)
summary["A6"], summary["B6"] = "Official-feed matches", keep_groups
summary["A7"], summary["B7"] = "Manual-review groups", manual_groups
summary["A8"], summary["B8"] = "Potential removals after merge", remove_candidates
summary["A9"], summary["B9"] = "Production changes made", "None"
for row_number in range(4, 10):
    for cell in summary[row_number]:
        cell.font = Font(name="Arial", size=10, bold=True, color=dark)
summary["A11"] = "Decision rules"
summary["A11"].font = Font(name="Arial", size=12, bold=True, color=dark)
rules = [
    "Use current Colorado MED and Michigan CRA active-license feeds as authoritative license evidence.",
    "Keep the record that best matches the official business name and location while protecting claimed and verified state.",
    "Merge complementary contact and profile fields before removing a duplicate record.",
    "Do not remove records marked Manual review. Confirm inactive, transferred, or mismatched licenses first.",
]
for index, rule in enumerate(rules, 12):
    summary.cell(index, 1, rule)
    summary.cell(index, 1).font = Font(name="Arial", size=10, color="18212B")
summary.column_dimensions["A"].width = 92
summary.column_dimensions["B"].width = 18

title(groups_sheet, "License-level decisions")
group_headers = ["License number", "Records", "Decision", "Canonical record ID", "Official business", "Official address", "Official status", "Official expiration", "Source provider", "Source URL"]
write_header(groups_sheet, 5, group_headers)
for row_number, (license, group) in enumerate(sorted(groups.items()), 6):
    first = group[0]
    values = [license, len(group), first["decision"], first["canonicalRecordId"], first["officialBusinessName"], first["officialAddress"], first["officialStatus"], first["officialExpiration"], first["sourceProvider"], first["sourceUrl"]]
    for column, value in enumerate(values, 1):
        groups_sheet.cell(row_number, column, value)
    if first["sourceUrl"]:
        groups_sheet.cell(row_number, 10).hyperlink = first["sourceUrl"]
        groups_sheet.cell(row_number, 10).style = "Hyperlink"
group_last = len(groups) + 5
style_body(groups_sheet, 6, group_last, len(group_headers))
groups_sheet.auto_filter.ref = f"A5:J{group_last}"
groups_sheet.freeze_panes = "D6"
groups_sheet.conditional_formatting.add(f"C6:C{group_last}", FormulaRule(formula=["LEFT(C6,6)=\"MANUAL\""], fill=PatternFill("solid", fgColor=amber)))
groups_sheet.conditional_formatting.add(f"C6:C{group_last}", FormulaRule(formula=["C6=\"KEEP AND MERGE DUPLICATES\""], fill=PatternFill("solid", fgColor=light_green)))
for column, width in enumerate([18, 10, 42, 25, 36, 48, 24, 18, 20, 48], 1):
    groups_sheet.column_dimensions[groups_sheet.cell(5, column).column_letter].width = width

title(records_sheet, "Record-level recommendations")
record_headers = ["License number", "Decision", "Recommendation", "Canonical record ID", "Record ID", "Business name", "DBA", "Address", "City", "State", "Postal code", "Claimed", "Canoja verified", "Official business", "Official address", "Official status", "Name match", "Street match", "City match", "ZIP match", "Match score", "Source URL"]
write_header(records_sheet, 5, record_headers)
for row_number, row in enumerate(rows, 6):
    values = [row["license"], row["decision"], row["recommendation"], row["canonicalRecordId"], row["recordId"], row["businessName"], row["dba"], row["address"], row["city"], row["state"], row["postalCode"], row["claimed"], row["canojaVerified"], row["officialBusinessName"], row["officialAddress"], row["officialStatus"], row["nameMatch"], row["streetMatch"], row["cityMatch"], row["zipMatch"], row["matchScore"], row["sourceUrl"]]
    for column, value in enumerate(values, 1):
        records_sheet.cell(row_number, column, value)
    if row["sourceUrl"]:
        records_sheet.cell(row_number, 22).hyperlink = row["sourceUrl"]
        records_sheet.cell(row_number, 22).style = "Hyperlink"
record_last = len(rows) + 5
style_body(records_sheet, 6, record_last, len(record_headers))
records_sheet.auto_filter.ref = f"A5:V{record_last}"
records_sheet.freeze_panes = "F6"
records_sheet.conditional_formatting.add(f"C6:C{record_last}", FormulaRule(formula=["C6=\"KEEP AS CANONICAL\""], fill=PatternFill("solid", fgColor=light_green)))
records_sheet.conditional_formatting.add(f"C6:C{record_last}", FormulaRule(formula=["C6=\"MANUAL REVIEW\""], fill=PatternFill("solid", fgColor=amber)))
records_sheet.conditional_formatting.add(f"C6:C{record_last}", FormulaRule(formula=["LEFT(C6,5)=\"MERGE\""], fill=PatternFill("solid", fgColor=red)))
for column, width in enumerate([18, 42, 34, 25, 25, 34, 30, 44, 20, 16, 12, 10, 16, 34, 44, 24, 12, 12, 12, 12, 12, 48], 1):
    records_sheet.column_dimensions[records_sheet.cell(5, column).column_letter].width = width

title(sources_sheet, "Official sources")
source_headers = ["Jurisdiction", "Publisher", "Title", "Accessed", "URL", "Use"]
write_header(sources_sheet, 5, source_headers)
source_rows = [
    ["Colorado", "Colorado Marijuana Enforcement Division", "MED Licensed Facilities", datetime.now(), "https://med.colorado.gov/licensee-information-and-lookup-tool/licensed-facilities", "Current active facility name, address, type, and expiration"],
    ["Michigan", "Michigan Cannabis Regulatory Agency", "Verify a License", datetime.now(), "https://www.michigan.gov/cra/verify-a-license-1", "Current active retail license name, address, type, status, and expiration"],
]
for row_number, values in enumerate(source_rows, 6):
    for column, value in enumerate(values, 1):
        sources_sheet.cell(row_number, column, value)
    sources_sheet.cell(row_number, 5).hyperlink = values[4]
    sources_sheet.cell(row_number, 5).style = "Hyperlink"
    sources_sheet.cell(row_number, 4).number_format = "mm/dd/yyyy"
style_body(sources_sheet, 6, 7, len(source_headers))
for column, width in enumerate([16, 38, 28, 14, 58, 60], 1):
    sources_sheet.column_dimensions[sources_sheet.cell(5, column).column_letter].width = width

for sheet in wb.worksheets:
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "5:5" if sheet.title != "Summary" else None

output_path = base_dir / "duplicate_license_cleanup_recommendations.xlsx"
wb.save(output_path)
check = load_workbook(output_path, read_only=True, data_only=False)
assert check["Group decisions"].max_row == group_last
assert check["Record recommendations"].max_row == record_last
assert check["Summary"]["B4"].value == 225
assert check["Summary"]["B9"].value == "None"
print(json.dumps({"path": str(output_path), "groups": len(groups), "keepAndMerge": keep_groups, "manualReview": manual_groups, "potentialRemovals": remove_candidates}))
